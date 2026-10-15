from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.services.chunking import ChunkDraft, PageAwareChunker
from app.services.pdf_ingestion import ContextualVisionDescriber, VisualAsset


class DocxIngestionError(ValueError):
    pass


@dataclass(frozen=True)
class DocxExtraction:
    chunks: list[ChunkDraft]
    visuals: list[VisualAsset]


class DocxExtractor:
    """DOCX structure reader; derived assets stay in private storage."""

    def __init__(self, chunker: PageAwareChunker | None = None) -> None:
        self._chunker = chunker or PageAwareChunker()
        self._vision = ContextualVisionDescriber()

    def extract(self, data: bytes, *, source_id: str) -> DocxExtraction:
        if not data:
            raise DocxIngestionError("The DOCX file is empty.")
        try:
            document = Document(BytesIO(data))
        except Exception as exc:
            raise DocxIngestionError("The file is not a valid DOCX document.") from exc
        chunks: list[ChunkDraft] = []
        section: str | None = None
        nearby: list[str] = []
        for position, paragraph in enumerate(document.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name.lower() if paragraph.style else ""
            if style.startswith("heading"):
                section = text
                chunks.append(ChunkDraft(text, "docx_heading", section_title=section, metadata={"source_id": source_id, "extraction_method": "python-docx", "document_position": position, "heading_level": style}))
                continue
            nearby.append(text)
            content_type = "docx_list" if "list" in style else "docx_text"
            chunks.extend(self._chunker.text_chunks(text, content_type=content_type, section_title=section, metadata={"source_id": source_id, "extraction_method": "python-docx", "style": style, "document_position": position}))
        for table_index, table in enumerate(document.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            text = self._table_text(rows)
            if text:
                chunks.append(ChunkDraft(text, "docx_table", section_title=section, metadata={"source_id": source_id, "extraction_method": "python-docx", "table_index": table_index, "row_range": [1, len(rows)], "row_count": len(rows), "nearby_text": context if 'context' in locals() else ""}))
        visuals: list[VisualAsset] = []
        context = " ".join(nearby[-4:])[:500]
        for relation in document.part.rels.values():
            if relation.reltype != RT.IMAGE:
                continue
            image = relation.target_part
            extension = image.partname.ext.lstrip(".") or "png"
            if len(image.blob) >= 2048:
                visuals.append(VisualAsset(0, image.blob, extension, self._kind(context), context))
        if not chunks and not visuals:
            raise DocxIngestionError("No usable content was found in the DOCX document.")
        return DocxExtraction(chunks, visuals)

    def visual_chunk(self, asset: VisualAsset, asset_path: str, source_id: str) -> ChunkDraft:
        return ChunkDraft("Visual asset retained for configured vision analysis.", asset.kind, section_title=None, metadata={"source_id": source_id, "asset_path": asset_path, "extraction_method": "python-docx", "nearby_text": asset.nearby_text[:500], "visual_analysis_status": "pending_provider"})

    @staticmethod
    def _kind(context: str) -> str:
        lowered = context.lower()
        if "chart" in lowered or "graph" in lowered:
            return "docx_chart"
        if "diagram" in lowered or "figure" in lowered:
            return "docx_diagram"
        return "docx_image"

    @staticmethod
    def _table_text(rows: list[list[str]]) -> str:
        if not rows:
            return ""
        header = rows[0]
        lines = ["Table: " + " | ".join(header)]
        for number, row in enumerate(rows[1:], 1):
            pairs = [f"{header[index] or f'Column {index + 1}'}: {value}" for index, value in enumerate(row) if value]
            if pairs:
                lines.append(f"Row {number}: " + "; ".join(pairs))
        return "\n".join(lines) if len(lines) > 1 else ""
