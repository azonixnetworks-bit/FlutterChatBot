from io import BytesIO

import pytest
from docx import Document

from app.services.docx_ingestion import DocxExtractor, DocxIngestionError


def _document() -> bytes:
    document = Document()
    document.add_heading("Market Notes", level=1)
    document.add_paragraph("EURUSD moved into the demand zone.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Pair"
    table.rows[0].cells[1].text = "Bias"
    table.rows[1].cells[0].text = "EURUSD"
    table.rows[1].cells[1].text = "Bullish"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_extracts_heading_text_and_table() -> None:
    result = DocxExtractor().extract(_document(), source_id="source")
    assert any(chunk.content_type == "docx_text" and chunk.section_title == "Market Notes" for chunk in result.chunks)
    table = next(chunk for chunk in result.chunks if chunk.content_type == "docx_table")
    assert "Pair: EURUSD" in table.content


def test_preserves_headings_lists_and_document_order() -> None:
    document = Document()
    document.add_heading("Demand Zone Validation", level=1)
    document.add_paragraph("See the example below.")
    document.add_paragraph("First confirmation", style="List Bullet")
    output = BytesIO(); document.save(output)
    chunks = DocxExtractor().extract(output.getvalue(), source_id="source").chunks
    assert [chunk.content_type for chunk in chunks] == ["docx_heading", "docx_text", "docx_list"]
    assert chunks[-1].metadata["document_position"] > chunks[1].metadata["document_position"]


@pytest.mark.parametrize("data", [b"", b"not a docx"])
def test_rejects_empty_or_corrupted_docx(data: bytes) -> None:
    with pytest.raises(DocxIngestionError):
        DocxExtractor().extract(data, source_id="source")
